from __future__ import annotations

import pandas as pd
import pytest

from erl.export import (
    collect_results,
    export_results,
    headline_summary,
    sheet_name,
    write_text_summary,
    write_workbook,
)


def _fixture(directory) -> None:
    """A processed directory resembling a real run."""
    pd.DataFrame(
        {"rel_day": [-2, -1, 0, 1, 2], "mean_abs_ar": [0.011, 0.011, 0.021, 0.036, 0.014],
         "n": [2237, 2238, 2240, 2240, 2239]}
    ).to_csv(directory / "alignment_diagnostic.csv", index=False)
    pd.DataFrame(
        {"term": ["intercept", "treatment_mean", "runup_20d", "mcap_decile"],
         "coef": [-0.00005, 0.00505, -0.00055, 0.00199],
         "se": [0.00181, 0.00133, 0.00116, 0.00093],
         "tstat": [-0.03, 3.81, -0.47, 2.14]}
    ).to_csv(directory / "forest_blp.csv", index=False)
    pd.DataFrame(
        {"term": ["intercept", "runup_20d", "mcap_decile"],
         "coef": [0.00496, -0.00016, 0.00087],
         "se": [0.00008, 0.0000335, 0.0000947],
         "tstat": [60.4, -4.91, 9.20]}
    ).to_csv(directory / "forest_blp_naive.csv", index=False)
    pd.DataFrame(
        [{"wald_statistic": 9.21, "wald_dof": 3, "wald_pvalue": 0.0267, "n": 2115,
          "breaks": "2009-07-01;2020-03-01"}]
    ).to_csv(directory / "regime_stability_test.csv", index=False)
    pd.DataFrame(
        [{"wald_statistic": 5.56, "wald_dof": 3, "wald_pvalue": 0.135, "n": 2115,
          "target": "car_reaction_vol_adj"}]
    ).to_csv(directory / "regime_stability_voladj_test.csv", index=False)
    pd.DataFrame(
        {"group": ["still a member", "left the index"], "tickers": [478, 465],
         "harvested": [478, 355], "coverage": [1.0, 0.763]}
    ).to_csv(directory / "universe_coverage.csv", index=False)
    pd.DataFrame(
        {"model": ["ols", "lightgbm"], "r2": [0.0264, 0.0286],
         "mae": [0.0590, 0.0584], "rank_ic": [0.1475, 0.1714]}
    ).to_csv(directory / "prediction_comparison.csv", index=False)
    pd.DataFrame(
        [{"model_a": "lightgbm", "model_b": "ols", "mse_pvalue": 0.815, "ic_diff": 0.0239,
          "ic_ci_low": -0.0397, "ic_ci_high": 0.0864, "n": 423,
          "verdict": "indistinguishable"}]
    ).to_csv(directory / "prediction_significance.csv", index=False)
    pd.DataFrame(
        {"feature": ["sue", "runup_60d", "pe_z"], "n_splits": [240, 160, 0],
         "gain": [80239.0, 44062.0, 0.0], "gain_share": [0.263, 0.144, 0.0]}
    ).to_csv(directory / "gbm_feature_usage.csv", index=False)
    pd.DataFrame(
        {"year": [2006, 2022], "events": [102, 124], "placeholder": [13, 0],
         "placeholder_share": [0.127, 0.0]}
    ).to_csv(directory / "estimate_backfill_by_year.csv", index=False)


def test_collect_orders_headline_tables_first(tmp_path):
    _fixture(tmp_path)
    (tmp_path / "zzz_extra.csv").write_text("a,b\n1,2\n")
    tables = collect_results(tmp_path)
    keys = list(tables)
    assert keys[0] == "forest_blp"          # preferred order wins over alphabetical
    assert keys[-1] == "zzz_extra"          # unlisted files still included, at the end
    assert len(tables) == 11


def test_summary_extracts_the_numbers_that_matter(tmp_path):
    _fixture(tmp_path)
    summary = headline_summary(collect_results(tmp_path))
    text = summary.to_string()
    # alignment peak identified
    peak = summary.loc[summary["metric"].str.contains("peak at rel_day"), "value"].iloc[0]
    assert peak == 1
    # the honest-vs-naive contrast is surfaced, not buried
    assert "Largest SE understatement" in text
    ratio = summary.loc[
        summary["metric"].str.contains("understatement"), "value"
    ].iloc[0]
    assert float(str(ratio).rstrip("x")) > 30
    # only the genuinely significant moderator is listed
    sig = summary.loc[
        summary["metric"] == "Moderators significant at 5% (honest SEs)", "value"
    ].iloc[0]
    assert sig == "mcap_decile"
    # both stability verdicts present, with opposite conclusions
    raw = summary.loc[summary["metric"].str.contains(r"\(raw reaction\)"), "note"].iloc[0]
    adj = summary.loc[summary["metric"].str.contains(r"\(vol-adjusted reaction\)"), "note"].iloc[0]
    assert "rejects stability" in raw
    assert "does not reject" in adj
    # survivorship coverage is reported with its caveat
    note = summary.loc[
        summary["metric"].str.contains("left the index"), "note"
    ].iloc[0]
    assert "355 of 465" in note
    # every row names its source file
    assert summary["source"].str.contains(".csv").all()


def test_workbook_has_a_sheet_per_table_plus_summary(tmp_path):
    from openpyxl import load_workbook

    _fixture(tmp_path)
    tables = collect_results(tmp_path)
    summary = headline_summary(tables)
    path = write_workbook(tables, tmp_path / "results.xlsx", summary)
    book = load_workbook(path)
    assert book.sheetnames[0] == "Summary"
    assert set(tables) <= set(book.sheetnames)
    sheet = book["forest_blp"]
    assert sheet["A1"].value == "term"
    assert sheet["A1"].font.bold
    assert sheet.freeze_panes == "A2"
    # values, not formulas: these are computed outputs, not a model
    assert all(
        not (isinstance(c.value, str) and c.value.startswith("="))
        for row in sheet.iter_rows() for c in row
    )


def test_long_and_invalid_sheet_names_are_handled():
    taken: set[str] = set()
    a = sheet_name("regime_stability_voladj_test_with_a_very_long_name", taken)
    taken.add(a)
    assert len(a) <= 31
    b = sheet_name("regime_stability_voladj_test_with_a_very_long_name", taken)
    assert b != a and len(b) <= 31
    assert sheet_name("bad[name]:with/chars", set()) == "bad_name__with_chars"


def test_text_summary_truncates_long_tables(tmp_path):
    _fixture(tmp_path)
    pd.DataFrame({"effect": range(200), "se": range(200)}).to_csv(
        tmp_path / "rolling_effect.csv", index=False
    )
    tables = collect_results(tmp_path)
    path = write_text_summary(tables, tmp_path / "results.md", headline_summary(tables))
    text = path.read_text()
    assert "# Results" in text
    assert "## Headline numbers" in text
    assert "200 rows; first 30 and last 10 shown" in text


def test_export_is_a_no_op_on_an_empty_directory(tmp_path):
    assert export_results(tmp_path) == []
    assert not (tmp_path / "results.xlsx").exists()


def test_export_writes_all_three_artifacts(tmp_path):
    _fixture(tmp_path)
    written = export_results(tmp_path)
    names = {p.name for p in written}
    assert names == {"results.xlsx", "results.md", "results_summary.csv"}
    assert all(p.exists() and p.stat().st_size > 0 for p in written)


def test_provenance_flags_a_mixed_run(tmp_path):
    """A full-universe harvest with a stale pilot analysis must be reported, not
    presented as one coherent set of results."""
    from erl.export import provenance_check

    _fixture(tmp_path)
    pd.DataFrame([{"stage": "harvest", "universe": "sp500", "tickers_requested": 943,
                   "events": 54158, "tickers_with_events": 733}]).to_csv(
        tmp_path / "harvest_manifest.csv", index=False)
    pd.DataFrame([{"stage": "panel", "events": 2415, "tickers": 31,
                   "date_min": "2005-01-11", "date_max": "2026-09-10"}]).to_csv(
        tmp_path / "panel_manifest.csv", index=False)
    tables = collect_results(tmp_path)
    check = provenance_check(tables)
    assert check["status"] == "MISMATCH"
    assert "733" in check["detail"] and "31" in check["detail"]
    summary = headline_summary(tables)
    assert summary["metric"].iloc[0] == "Provenance of the tables below"
    assert summary["value"].iloc[0] == "MISMATCH"


def test_provenance_passes_on_a_single_run(tmp_path):
    from erl.export import provenance_check

    _fixture(tmp_path)
    pd.DataFrame([{"stage": "harvest", "universe": "sp500", "tickers_requested": 943,
                   "events": 54158, "tickers_with_events": 733}]).to_csv(
        tmp_path / "harvest_manifest.csv", index=False)
    pd.DataFrame([{"stage": "panel", "events": 48000, "tickers": 700,
                   "date_min": "2005-01-03", "date_max": "2026-09-11"}]).to_csv(
        tmp_path / "panel_manifest.csv", index=False)
    assert provenance_check(collect_results(tmp_path))["status"] == "consistent"


def test_provenance_is_unknown_without_manifests(tmp_path):
    from erl.export import provenance_check

    _fixture(tmp_path)
    assert provenance_check(collect_results(tmp_path))["status"] == "unknown"
